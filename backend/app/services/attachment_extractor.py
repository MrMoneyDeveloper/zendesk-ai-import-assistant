from io import BytesIO
import mimetypes
from pathlib import Path

from fastapi import UploadFile

from app.core.settings import Settings

TEXT_EXTENSIONS = {
    ".txt",
    ".csv",
    ".json",
    ".md",
    ".log",
    ".yaml",
    ".yml",
}

TEXT_MIME_TYPES = {
    "text/plain",
    "text/csv",
    "application/json",
    "application/x-yaml",
    "text/yaml",
}

PDF_MIME_TYPES = {
    "application/pdf",
}

DOCX_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class AttachmentExtractionError(RuntimeError):
    pass


def _extract_pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise AttachmentExtractionError(
            "PDF extraction dependency is missing. Install `pypdf`."
        ) from exc
    reader = PdfReader(BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            parts.append(page_text.strip())
    return "\n\n".join(parts)


def _extract_docx_text(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise AttachmentExtractionError(
            "DOCX extraction dependency is missing. Install `python-docx`."
        ) from exc
    doc = Document(BytesIO(data))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_values:
                parts.append(" | ".join(row_values))
    return "\n".join(parts)


def _decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("utf-8", errors="ignore")


async def extract_attachment_payload(file: UploadFile, settings: Settings) -> dict:
    filename = (file.filename or "attachment").strip() or "attachment"
    extension = Path(filename).suffix.lower()
    mime_type = (file.content_type or "").strip().lower()
    guessed_mime, _ = mimetypes.guess_type(filename)
    if not mime_type:
        mime_type = (guessed_mime or "").lower()

    data = await file.read()
    if len(data) > settings.attachment_max_file_size_bytes:
        raise AttachmentExtractionError(
            "Attachment exceeds max size "
            f"({settings.attachment_max_file_size_bytes} bytes)."
        )

    warnings: list[str] = []
    extracted_text = ""

    if extension in TEXT_EXTENSIONS or mime_type.startswith("text/") or mime_type in TEXT_MIME_TYPES:
        extracted_text = _decode_text(data)
    elif extension == ".pdf" or mime_type in PDF_MIME_TYPES:
        extracted_text = _extract_pdf_text(data)
    elif extension == ".docx" or mime_type in DOCX_MIME_TYPES:
        extracted_text = _extract_docx_text(data)
    else:
        raise AttachmentExtractionError(
            "Unsupported attachment type. Allowed: txt, csv, json, md, log, yaml, yml, pdf, docx."
        )

    extracted_text = extracted_text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not extracted_text:
        warnings.append("No readable text was extracted from the attachment.")

    truncated = False
    if len(extracted_text) > settings.attachment_max_chars:
        extracted_text = extracted_text[: settings.attachment_max_chars]
        truncated = True
        warnings.append(
            f"Text was truncated to {settings.attachment_max_chars} characters."
        )

    return {
        "filename": filename,
        "mime_type": mime_type or (guessed_mime or "application/octet-stream"),
        "char_count": len(extracted_text),
        "extracted_text": extracted_text,
        "truncated": truncated,
        "warnings": warnings,
    }
